import streamlit as st
import pandas as pd
import openpyxl
from datetime import datetime
import io
import re

# 页面基础配置
st.set_page_config(
    page_title="OOH 投放结案 Summary 自动生成器",
    page_icon="📊",
    layout="wide"
)

st.title("📊 OOH 投放结案 Summary 自动化生成工具")
st.write("请依次上传 **1《统计DB》**、**2《Spotplan》**、**3《补偿汇总》（可选）** 以及 **4《OOH投放结案Summary格式》**。系统将以格式表为蓝本，自动填入数据并完美保留原始样式、颜色格式和 Excel 动态计算公式[cite: 5]。")

st.divider()

# 文件上传区域 (4列布局，将补偿汇总改为可选)
col1, col2, col3, col4 = st.columns(4)

with col1:
    db_file = st.file_uploader("1、上传《统计DB》", type=["xlsx"], key="db")

with col2:
    spot_file = st.file_uploader("2、上传《Spotplan》", type=["xlsx"], key="spot")

with col3:
    comp_file = st.file_uploader("3、上传《补偿汇总》(可选)", type=["xlsx", "xls"], key="comp")

with col4:
    template_file = st.file_uploader("4、上传《OOH投放结案Summary格式》", type=["xlsx"], key="template")

def clean_location_name(loc_name):
    """提取并清洗媒体名称，去掉（赠送/额外赠送/增值）字眼"""
    if not loc_name:
        return "", "", ""
    loc = str(loc_name).strip()
    loc_clean = re.sub(r'[（\(](赠送|额外赠送|增值)[）\)]', '', loc).strip()
    loc_clean2 = re.sub(r'[（\(]\d+块/套[）\)]', '', loc_clean).strip()
    return loc_clean, loc_clean2, loc

def generate_summary_from_template(spot_file, db_file, comp_file, template_file):
    # -------------------------------------------------------------
    # 1. 读取 统计 DB 构建 Lookup 检索字典[cite: 5]
    # -------------------------------------------------------------
    wb_db = openpyxl.load_workbook(db_file, data_only=True)
    ws_db = wb_db.active
    if '统计DB' in wb_db.sheetnames:
        ws_db = wb_db['统计DB']
    
    db_lookup = {}
    for r in range(2, ws_db.max_row + 1):
        loc_val = ws_db.cell(r, 13).value       # Col M: Location
        coverage_val = ws_db.cell(r, 40).value  # Col AN: 有效覆盖人车次/天（单块）
        
        if loc_val is not None:
            loc_str = str(loc_val).strip()
            if loc_str and loc_str not in db_lookup:
                db_lookup[loc_str] = coverage_val
            c1, c2, _ = clean_location_name(loc_str)
            if c1 and c1 not in db_lookup:
                db_lookup[c1] = coverage_val
            if c2 and c2 not in db_lookup:
                db_lookup[c2] = coverage_val

    # -------------------------------------------------------------
    # 2. 读取 赔付统计表 构建 Compensation 检索字典（支持可选跳过）[cite: 5]
    # -------------------------------------------------------------
    comp_lookup = {}
    if comp_file is not None:
        wb_comp = openpyxl.load_workbook(comp_file, data_only=True)
        ws_comp = wb_comp.active
        
        comp_header_row = 1
        for r in range(1, 10):
            row_vals = [str(ws_comp.cell(r, c).value or '') for c in range(1, 15)]
            if any('媒体形式' in v for v in row_vals):
                comp_header_row = r
                break

        header_map = {}
        for c in range(1, ws_comp.max_column + 1):
            val = str(ws_comp.cell(comp_header_row, c).value or '').strip()
            if val:
                header_map[val] = c

        col_city = header_map.get('城市')
        col_media = header_map.get('媒体形式')
        col_anomaly = header_map.get('异常情况')
        col_plan = header_map.get('补偿方案')
        col_amount = header_map.get('实际补偿(元)') or header_map.get('应补偿价值(元)') or header_map.get('受影响价值(元)')

        if col_media:
            for r in range(comp_header_row + 1, ws_comp.max_row + 1):
                media_val = ws_comp.cell(r, col_media).value
                if not media_val:
                    continue
                    
                city_val = str(ws_comp.cell(r, col_city).value or '').strip() if col_city else ""
                media_str = str(media_val).strip()
                anomaly_val = ws_comp.cell(r, col_anomaly).value if col_anomaly else ""
                plan_val = ws_comp.cell(r, col_plan).value if col_plan else ""
                amount_val = ws_comp.cell(r, col_amount).value if col_amount else 0

                comp_info = {
                    'anomaly': str(anomaly_val or '').strip(),
                    'plan': str(plan_val or '').strip(),
                    'amount': amount_val if amount_val is not None else 0
                }

                comp_lookup[media_str] = comp_info
                if city_val and not media_str.startswith(city_val):
                    comp_lookup[f"{city_val}{media_str}"] = comp_info
                
                c1, c2, _ = clean_location_name(media_str)
                if c1:
                    comp_lookup[c1] = comp_info
                if c2:
                    comp_lookup[c2] = comp_info

    # -------------------------------------------------------------
    # 3. 读取 Spotplan 明细数据[cite: 5]
    # -------------------------------------------------------------
    wb_spot = openpyxl.load_workbook(spot_file, data_only=True)
    ws_spot = wb_spot.active
    
    header_row = 4
    for r in range(1, 10):
        row_vals = [str(ws_spot.cell(r, c).value or '') for c in range(1, 15)]
        if 'Market' in row_vals and 'Location' in row_vals:
            header_row = r
            break
            
    spot_rows = []
    for r in range(header_row + 1, ws_spot.max_row + 1):
        mkt = ws_spot.cell(r, 2).value  # Col B
        loc = ws_spot.cell(r, 4).value  # Col D
        if mkt or loc:
            row_data = [ws_spot.cell(r, c).value for c in range(2, 18)]
            spot_rows.append(row_data)

    # -------------------------------------------------------------
    # 4. 加载 Summary 模板文件（保留原始样式格式与表格头结构）[cite: 5]
    # -------------------------------------------------------------
    wb_tpl = openpyxl.load_workbook(template_file)
    ws_tpl = wb_tpl.active
    
    start_row = 4  # 数据写入起始行
    
    # 解除数据写入区域的合并单元格（避免赋值冲突）
    merged_ranges = list(ws_tpl.merged_cells.ranges)
    for rng in merged_ranges:
        if rng.max_row >= start_row:
            ws_tpl.unmerge_cells(str(rng))

    # 映射主媒体与其对应的“行号”及“折扣值”
    main_media_info = {} 
    for idx, row in enumerate(spot_rows):
        current_row = start_row + idx
        loc = str(row[2] or '')
        c1, c2, orig = clean_location_name(loc)
        is_bonus = ("赠送" in orig) or ("额外赠送" in orig) or ("增值" in orig)
        
        if not is_bonus and c1:
            discount_val = row[10] # Col L: Discount
            main_media_info[c1] = {'row': current_row, 'discount': discount_val}
            if c2:
                main_media_info[c2] = {'row': current_row, 'discount': discount_val}

    sample_cells = [ws_tpl.cell(start_row, col) for col in range(1, 32)]

    # -------------------------------------------------------------
    # 5. 填充数据并嵌入原生 Excel 计算公式与格式[cite: 5]
    # -------------------------------------------------------------
    for idx, row in enumerate(spot_rows):
        current_row = start_row + idx
        
        mkt = row[0]          
        fmt = row[1]          
        loc = str(row[2] or '')
        period = row[3]       
        no_week = row[4]      
        no_unit = row[5]      
        buying_unit = row[6]  
        duration_freq = row[7]
        ratecard_cost = row[8]
        ratecard_ttl = row[9] 
        discount = row[10]    
        net_unit_cost = row[11]
        net_ttl_cost = row[12] 
        unit_prod_fee = row[13]
        prod_fee = row[14]    
        gross_cost = row[15]  

        c1, c2, orig = clean_location_name(loc)
        is_bonus = ("赠送" in orig) or ("额外赠送" in orig) or ("增值" in orig)

        matched_comp = comp_lookup.get(orig) or comp_lookup.get(c1) or comp_lookup.get(c2) or comp_lookup.get(f"{mkt}{c1}")
        
        comp_anomaly = matched_comp['anomaly'] if matched_comp else ""
        comp_plan = matched_comp['plan'] if matched_comp else ""
        comp_amount = matched_comp['amount'] if matched_comp else 0

        y_net_value_formula = 0
        if is_bonus:
            matched_main = main_media_info.get(c1) or main_media_info.get(c2)
            if matched_main:
                main_row = matched_main['row']
                y_net_value_formula = f"=K{current_row}*L{main_row}"
            else:
                y_net_value_formula = f"=K{current_row}*L{current_row}"

        resource_qty = f"{no_unit}{buying_unit}" if (no_unit and buying_unit) else no_unit

        daily_coverage = db_lookup.get(orig) or db_lookup.get(c1) or db_lookup.get(c2)
        if daily_coverage == "/" or daily_coverage is None:
            for k, v in db_lookup.items():
                if c2 and c2 in k and v != "/":
                    daily_coverage = v
                    break

        row_values = {
            1: mkt,                                                        
            2: loc,                                                        
            3: resource_qty,                                               
            4: duration_freq,                                              
            5: period,                                                     
            6: f"=(_xlfn.TEXTAFTER(E{current_row},\"-\")-_xlfn.TEXTBEFORE(E{current_row},\"-\")+1)/7", 
            7: no_unit,                                                    
            8: buying_unit,                                                
            9: duration_freq,                                              
            10: ratecard_cost,                                             
            11: f"=J{current_row}*G{current_row}*F{current_row}",          
            12: discount if not is_bonus else 0,                           
            13: f"=J{current_row}*L{current_row}",                          
            14: f"=ROUND(M{current_row}*F{current_row}*G{current_row},0)", 
            15: unit_prod_fee if not is_bonus else 0,                      
            16: f"=O{current_row}*G{current_row}",                          
            17: f"=N{current_row}+P{current_row}",                          
            18: comp_anomaly,                                              
            19: comp_plan,                                                 
            20: 0 if not is_bonus else "",                                 
            21: f"=_xlfn.TEXTBEFORE(E{current_row},\"-\")&\"-\"&_xlfn.TEXTAFTER(E{current_row+1},\"-\")" if not is_bonus else "", 
            22: f"=N{current_row}" if not is_bonus else 0,                  
            23: f"=P{current_row}" if not is_bonus else 0,                  
            24: 0 if not is_bonus else f"=K{current_row}",                  
            25: 0 if not is_bonus else y_net_value_formula,                
            26: comp_amount,                                               
            27: 0,                                                         
            28: 0,                                                         
            29: daily_coverage if (daily_coverage is not None and not is_bonus) else "", 
            30: f"=_xlfn.TEXTAFTER(U{current_row},\"-\")-_xlfn.TEXTBEFORE(U{current_row},\"-\")+1" if not is_bonus else "", 
            31: f"=AC{current_row}*AD{current_row}" if not is_bonus else "" 
        }

        for col_idx, val in row_values.items():
            cell = ws_tpl.cell(row=current_row, column=col_idx)
            cell.value = val
            
            if col_idx - 1 < len(sample_cells):
                sample_cell = sample_cells[col_idx - 1]
                if sample_cell.has_style:
                    cell.font = sample_cell.font.copy()
                    cell.fill = sample_cell.fill.copy()
                    cell.border = sample_cell.border.copy()
                    cell.alignment = sample_cell.alignment.copy()
                    cell.number_format = sample_cell.number_format

    output = io.BytesIO()
    wb_tpl.save(output)
    output.seek(0)
    return output

# 触发按钮逻辑（仅需核心的 统计DB、Spotplan、模板 即可运行，补偿汇总可选）
if spot_file and db_file and template_file:
    if st.button("🚀 套用模板生成结案 Summary", type="primary"):
        try:
            with st.spinner("正在匹配数据、保留模板格式并写入 Excel 公式..."):
                excel_out = generate_summary_from_template(spot_file, db_file, comp_file, template_file)
                st.success("🎉 生成成功！数据已完全复刻蓝本格式，且 Excel 公式均已嵌套。")
                st.download_button(
                    label="📥 点击下载生成的结案 Summary.xlsx",
                    data=excel_out,
                    file_name="OOH_投放结案_Summary_已生成.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
        except Exception as e:
            st.error(f"生成失败，错误原因: {str(e)}")
else:
    st.info("💡 请在上方的《统计DB》、《Spotplan》和《OOH投放结案Summary格式》中上传核心文件以启用生成按钮（《补偿汇总》可根据项目实际情况选择性上传）。")